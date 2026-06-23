from __future__ import annotations

import hashlib
import math
from pathlib import Path

import numpy as np
import openpyxl
import scipy.io
import xlsxwriter
from xlsxwriter.utility import xl_col_to_name


OUTPUT_DIR_NAME = "excel files version"
NUMERIC_FIELDS = ["time", "relativeTime", "voltage", "current", "temperature"]
VALUE_CHUNK_SIZE = 1000


def scalar_to_text(value: object) -> str:
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float):
        return repr(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, bool):
        return "True" if value else "False"
    if isinstance(value, str):
        return value
    raise TypeError(f"Unsupported scalar type: {type(value)!r}")


def value_to_text_list(value: object) -> list[str]:
    if isinstance(value, np.ndarray):
        return [scalar_to_text(item) for item in value.reshape(-1)]
    return [scalar_to_text(value)]


def value_length(value: object) -> int:
    if isinstance(value, np.ndarray):
        return int(value.size)
    return 1


def chunk_values(text_values: list[str]) -> list[str]:
    return [
        "|".join(text_values[index : index + VALUE_CHUNK_SIZE])
        for index in range(0, len(text_values), VALUE_CHUNK_SIZE)
    ]


def update_hash(digest: "hashlib._Hash", values: list[str]) -> None:
    for value in values:
        encoded = value.encode("utf-8")
        digest.update(len(encoded).to_bytes(8, "big"))
        digest.update(encoded)
    digest.update(b"\xff")


def write_readme(worksheet: xlsxwriter.worksheet.Worksheet) -> None:
    lines = [
        "This workbook is an exact-value export of the source MAT file.",
        "Top-level text fields are written as plain text.",
        "Per-step numeric arrays are stored as pipe-delimited text chunks.",
        f"Each chunk cell contains up to {VALUE_CHUNK_SIZE} values.",
        "Numeric values are stored as text so Excel does not round or alter them.",
        "Use the step_view sheet to reconstruct one MATLAB step as a normal table.",
        "In step_view!B1, enter the step_index you want to inspect.",
        "The validation sheet contains SHA-256 hashes generated from the MAT source.",
        "Those hashes are verified again after the workbook is written.",
    ]
    worksheet.write_row(0, 0, ["note"])
    for row_index, line in enumerate(lines, start=1):
        worksheet.write_string(row_index, 0, line)


def write_step_view(
    workbook: xlsxwriter.Workbook,
    max_chunks: int,
    total_steps: int,
) -> None:
    worksheet = workbook.add_worksheet("step_view")
    worksheet.freeze_panes(7, 0)
    worksheet.set_column(0, 0, 18)
    worksheet.set_column(1, 6, 22)

    chunk_start_col = xl_col_to_name(2, col_abs=True)
    chunk_end_col = xl_col_to_name(1 + max_chunks, col_abs=True)

    worksheet.write_string(0, 0, "selected_step_index")
    worksheet.write_number(0, 1, 1)
    worksheet.data_validation(
        0,
        1,
        0,
        1,
        {
            "validate": "integer",
            "criteria": "between",
            "minimum": 1,
            "maximum": total_steps,
        },
    )

    worksheet.write_string(1, 0, "comment")
    worksheet.write_formula(1, 1, '=XLOOKUP($B$1,step_index!$A:$A,step_index!$B:$B,"")')
    worksheet.write_string(2, 0, "type")
    worksheet.write_formula(2, 1, '=XLOOKUP($B$1,step_index!$A:$A,step_index!$C:$C,"")')
    worksheet.write_string(3, 0, "date")
    worksheet.write_formula(3, 1, '=XLOOKUP($B$1,step_index!$A:$A,step_index!$D:$D,"")')
    worksheet.write_string(4, 0, "sample_count")
    worksheet.write_formula(4, 1, '=XLOOKUP($B$1,step_index!$A:$A,step_index!$E:$E,0)')
    worksheet.write_string(5, 0, "note")
    worksheet.write_string(
        5,
        1,
        "Change B1 to any step index. The table below spills the exact text values for that MATLAB step.",
    )

    worksheet.write_row(
        6,
        0,
        ["sample_index", "time", "relativeTime", "voltage", "current", "temperature"],
    )
    worksheet.write_formula(7, 0, "=SEQUENCE($B$5)")

    for column_index, field_name in enumerate(NUMERIC_FIELDS, start=1):
        formula = (
            f'=LET(n,$B$5,'
            f'm,MATCH($B$1,{field_name}!$A:$A,0),'
            f'chunkCount,ROUNDUP(n/{VALUE_CHUNK_SIZE},0),'
            f'chunks,INDEX({field_name}!{chunk_start_col}:{chunk_end_col},m,SEQUENCE(1,chunkCount)),'
            f'vals,TOCOL(TEXTSPLIT(TEXTJOIN("|",TRUE,chunks),"|"),1),'
            f'TAKE(vals,n))'
        )
        worksheet.write_formula(7, column_index, formula)


def export_workbook(mat_path: Path, output_root: Path) -> Path:
    relative_output = mat_path.with_suffix(".xlsx").name
    output_path = output_root / relative_output
    data = scipy.io.loadmat(mat_path, struct_as_record=False, squeeze_me=True)["data"]
    steps = data.step
    sample_counts = [value_length(step.time) for step in steps]
    max_chunks = max(math.ceil(max(sample_counts) / VALUE_CHUNK_SIZE), 1)
    hashes: dict[str, str] = {}

    if output_path.exists():
        output_path.unlink()

    workbook = xlsxwriter.Workbook(
        str(output_path),
        {"constant_memory": True, "strings_to_numbers": False},
    )
    workbook.use_zip64()

    readme_ws = workbook.add_worksheet("README")
    write_readme(readme_ws)

    metadata_ws = workbook.add_worksheet("metadata")
    metadata_ws.freeze_panes(1, 0)
    metadata_ws.write_row(0, 0, ["field", "value"])
    metadata_hash = hashlib.sha256()
    metadata_rows = [
        ["source_file", mat_path.name],
        ["procedure", scalar_to_text(data.procedure)],
        ["description", scalar_to_text(data.description)],
        ["total_steps", str(len(steps))],
        ["value_chunk_size", str(VALUE_CHUNK_SIZE)],
    ]
    for row_index, row_values in enumerate(metadata_rows, start=1):
        metadata_ws.write_string(row_index, 0, row_values[0])
        metadata_ws.write_string(row_index, 1, row_values[1])
        update_hash(metadata_hash, row_values)
    hashes["metadata"] = metadata_hash.hexdigest()

    step_index_ws = workbook.add_worksheet("step_index")
    step_index_ws.freeze_panes(1, 0)
    step_index_ws.write_row(0, 0, ["step_index", "comment", "type", "date", "sample_count"])
    step_index_hash = hashlib.sha256()
    for row_index, step in enumerate(steps, start=1):
        sample_count = sample_counts[row_index - 1]
        row_values = [
            str(row_index),
            scalar_to_text(step.comment),
            scalar_to_text(step.type),
            scalar_to_text(step.date),
            str(sample_count),
        ]
        step_index_ws.write_number(row_index, 0, row_index)
        step_index_ws.write_string(row_index, 1, row_values[1])
        step_index_ws.write_string(row_index, 2, row_values[2])
        step_index_ws.write_string(row_index, 3, row_values[3])
        step_index_ws.write_number(row_index, 4, sample_count)
        update_hash(step_index_hash, row_values)
    hashes["step_index"] = step_index_hash.hexdigest()

    for field_name in NUMERIC_FIELDS:
        worksheet = workbook.add_worksheet(field_name)
        worksheet.freeze_panes(1, 0)
        headers = ["step_index", "sample_count"]
        headers.extend(
            [
                f"values_{start:06d}_{min(start + VALUE_CHUNK_SIZE - 1, max_chunks * VALUE_CHUNK_SIZE):06d}"
                for start in range(1, max_chunks * VALUE_CHUNK_SIZE + 1, VALUE_CHUNK_SIZE)
            ]
        )
        worksheet.write_row(0, 0, headers)
        field_hash = hashlib.sha256()
        for row_index, step in enumerate(steps, start=1):
            sample_count = sample_counts[row_index - 1]
            text_values = value_to_text_list(getattr(step, field_name))
            chunks = chunk_values(text_values)
            row_values = [str(row_index), str(sample_count), *chunks]
            worksheet.write_number(row_index, 0, row_index)
            worksheet.write_number(row_index, 1, sample_count)
            for column_index, chunk in enumerate(chunks, start=2):
                worksheet.write_string(row_index, column_index, chunk)
            update_hash(field_hash, row_values)
        hashes[field_name] = field_hash.hexdigest()

    write_step_view(workbook, max_chunks=max_chunks, total_steps=len(steps))

    validation_ws = workbook.add_worksheet("validation")
    validation_ws.freeze_panes(1, 0)
    validation_ws.write_row(0, 0, ["sheet_name", "source_sha256"])
    for row_index, (sheet_name, digest) in enumerate(hashes.items(), start=1):
        validation_ws.write_string(row_index, 0, sheet_name)
        validation_ws.write_string(row_index, 1, digest)

    workbook.close()
    verify_workbook(output_path, hashes)
    return output_path


def verify_workbook(workbook_path: Path, expected_hashes: dict[str, str]) -> None:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        for sheet_name, expected_hash in expected_hashes.items():
            worksheet = workbook[sheet_name]
            digest = hashlib.sha256()
            for row_index, row in enumerate(worksheet.iter_rows(values_only=True)):
                if row_index == 0:
                    continue
                if sheet_name == "metadata":
                    values = ["" if row[0] is None else str(row[0]), "" if row[1] is None else str(row[1])]
                elif sheet_name == "step_index":
                    values = [
                        str(int(row[0])),
                        "" if row[1] is None else str(row[1]),
                        "" if row[2] is None else str(row[2]),
                        "" if row[3] is None else str(row[3]),
                        str(int(row[4])),
                    ]
                else:
                    sample_count = int(row[1])
                    chunk_count = max(math.ceil(sample_count / VALUE_CHUNK_SIZE), 1)
                    values = [str(int(row[0])), str(sample_count)]
                    values.extend("" if row[column] is None else str(row[column]) for column in range(2, 2 + chunk_count))
                update_hash(digest, values)
            actual_hash = digest.hexdigest()
            if actual_hash != expected_hash:
                raise ValueError(
                    f"Verification failed for {workbook_path.name} sheet {sheet_name}: "
                    f"{actual_hash} != {expected_hash}"
                )
    finally:
        workbook.close()


def main() -> None:
    root = Path.cwd()
    mat_files = sorted(root.glob("*.mat"))
    if not mat_files:
        raise FileNotFoundError("No .mat files were found in the current directory.")

    output_root = root / OUTPUT_DIR_NAME
    output_root.mkdir(exist_ok=True)

    for mat_path in mat_files:
        print(f"Exporting {mat_path.name}...")
        output_path = export_workbook(mat_path, output_root)
        print(f"Verified {output_path.name}")

    print(f"Created Excel exports in: {output_root}")


if __name__ == "__main__":
    main()
