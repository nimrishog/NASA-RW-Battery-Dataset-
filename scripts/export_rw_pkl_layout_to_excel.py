from __future__ import annotations

import math
import pickle
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import numpy as np
import openpyxl
import pandas as pd
import scipy.io
import xlsxwriter


OUTPUT_DIR_NAME = "excel files pkl layout"
ROWS_PER_SHEET = 1_048_575  # Excel limit minus header row
DATA_COLUMNS = ["Voltage", "Current", "time", "relTime", "cycle", "Capacity"]
RW9_PICKLE_NAME = "rw9_datacapa.pkl"


@dataclass
class ChargeStep:
    cycle: int
    mat_step_index: int
    date: str
    sample_count: int
    voltage: np.ndarray
    current: np.ndarray
    time: np.ndarray
    rel_time: np.ndarray


def load_mat_steps(mat_path: Path) -> list[ChargeStep]:
    data = scipy.io.loadmat(mat_path, struct_as_record=False, squeeze_me=True)["data"]
    steps = data.step
    charge_steps: list[ChargeStep] = []
    cycle = 0
    for mat_step_index, step in enumerate(steps, start=1):
        sample_count = int(step.time.size) if isinstance(step.time, np.ndarray) else 1
        if step.comment == "charge (random walk)" and step.type == "C" and sample_count > 1:
            cycle += 1
            charge_steps.append(
                ChargeStep(
                    cycle=cycle,
                    mat_step_index=mat_step_index,
                    date=str(step.date),
                    sample_count=sample_count,
                    voltage=np.asarray(step.voltage, dtype=np.float64).reshape(-1),
                    current=np.asarray(step.current, dtype=np.float64).reshape(-1),
                    time=np.asarray(step.time, dtype=np.float64).reshape(-1),
                    rel_time=np.asarray(step.relativeTime, dtype=np.float64).reshape(-1),
                )
            )
    return charge_steps


def float_to_text(value: float) -> str:
    return repr(float(value))


def build_rw9_capacity_map(mat_steps: list[ChargeStep], pickle_path: Path) -> list[float]:
    with pickle_path.open("rb") as handle:
        df = pickle.load(handle)
    if not isinstance(df, pd.DataFrame):
        raise TypeError(f"{pickle_path.name} does not contain a pandas DataFrame.")

    expected_rows = sum(step.sample_count for step in mat_steps)
    if len(df) != expected_rows:
        raise ValueError(f"RW9 pickle row count mismatch: {len(df)} != {expected_rows}")
    if int(df["cycle"].nunique()) != len(mat_steps):
        raise ValueError(
            f"RW9 pickle cycle count mismatch: {df['cycle'].nunique()} != {len(mat_steps)}"
        )

    expected_cycle = np.repeat(
        np.arange(1, len(mat_steps) + 1, dtype=np.int64),
        [step.sample_count for step in mat_steps],
    )
    actual_cycle = df["cycle"].to_numpy(dtype=np.int64)
    if not np.array_equal(expected_cycle, actual_cycle):
        raise ValueError("RW9 pickle 'cycle' column does not match the MAT-derived cycles.")

    for column_name, attribute_name in (
        ("Voltage", "voltage"),
        ("Current", "current"),
        ("time", "time"),
        ("relTime", "rel_time"),
    ):
        expected = np.concatenate([getattr(step, attribute_name) for step in mat_steps])
        actual = df[column_name].to_numpy(dtype=np.float64, copy=False)
        if not np.array_equal(expected, actual):
            mismatch_index = int(np.flatnonzero(expected != actual)[0])
            raise ValueError(
                f"RW9 pickle column '{column_name}' does not match MAT data at row {mismatch_index}."
            )

    capacity_by_cycle = (
        df.groupby("cycle", sort=True)["Capacity"]
        .first()
        .to_numpy(dtype=np.float64, copy=False)
    )
    if capacity_by_cycle.shape[0] != len(mat_steps):
        raise ValueError("RW9 capacity map length mismatch.")
    return capacity_by_cycle.tolist()


def compute_reference_discharge_rows(mat_path: Path, mat_steps: list[ChargeStep]) -> list[list[object]]:
    data = scipy.io.loadmat(mat_path, struct_as_record=False, squeeze_me=True)["data"]
    steps = data.step
    cycle_by_step_index = {step.mat_step_index: step.cycle for step in mat_steps}

    rows: list[list[object]] = []
    current_cycle = 0
    for mat_step_index, step in enumerate(steps, start=1):
        current_cycle = cycle_by_step_index.get(mat_step_index, current_cycle)
        if step.comment == "reference discharge" and step.type == "D":
            rel_time = np.asarray(step.relativeTime, dtype=np.float64).reshape(-1)
            current = np.asarray(step.current, dtype=np.float64).reshape(-1)
            capacity_ah = float(np.trapz(current, rel_time) / 3600.0)
            rows.append(
                [
                    mat_step_index,
                    str(step.date),
                    current_cycle,
                    capacity_ah,
                    int(rel_time.size),
                ]
            )
    return rows


def write_readme(
    worksheet: xlsxwriter.worksheet.Worksheet,
    mat_name: str,
    cycle_count: int,
    row_count: int,
    has_verified_capacity: bool,
) -> None:
    notes = [
        f"Source MAT file: {mat_name}",
        "This workbook uses the verified RW9 pickle-style row layout.",
        "Rows come from MATLAB steps where comment='charge (random walk)', type='C', and sample_count > 1.",
        "Each exported row is one sample from the selected MATLAB step.",
        "Floating-point values are stored as text in the data sheets so Excel does not rewrite the source values.",
        f"Total exported cycles: {cycle_count}",
        f"Total exported sample rows: {row_count}",
    ]
    if has_verified_capacity:
        notes.append("Capacity values were copied from the verified pickle after exact column-by-column validation against RW9.mat.")
    else:
        notes.append("Capacity is intentionally blank because no verified per-row capacity label file was provided for this battery.")
    notes.extend(
        [
            "The reference_discharge_capacity sheet lists measured reference-discharge capacities from the MAT file without guessing per-row labels.",
            "The cycle_map sheet links each exported cycle back to the original MATLAB step index and date.",
        ]
    )
    worksheet.write_row(0, 0, ["note"])
    for row_index, note in enumerate(notes, start=1):
        worksheet.write_string(row_index, 0, note)


def iter_data_rows(
    mat_steps: list[ChargeStep],
    capacity_by_cycle: list[float] | None,
) -> Iterable[list[object]]:
    for step in mat_steps:
        capacity = capacity_by_cycle[step.cycle - 1] if capacity_by_cycle is not None else None
        for sample_index in range(step.sample_count):
            yield [
                float_to_text(step.voltage[sample_index]),
                float_to_text(step.current[sample_index]),
                float_to_text(step.time[sample_index]),
                float_to_text(step.rel_time[sample_index]),
                step.cycle,
                float_to_text(capacity) if capacity is not None else None,
            ]


def write_cycle_map_sheet(
    workbook: xlsxwriter.Workbook, mat_steps: list[ChargeStep]
) -> None:
    worksheet = workbook.add_worksheet("cycle_map")
    worksheet.freeze_panes(1, 0)
    worksheet.write_row(0, 0, ["cycle", "mat_step_index", "date", "sample_count"])
    for row_index, step in enumerate(mat_steps, start=1):
        worksheet.write_number(row_index, 0, step.cycle)
        worksheet.write_number(row_index, 1, step.mat_step_index)
        worksheet.write_string(row_index, 2, step.date)
        worksheet.write_number(row_index, 3, step.sample_count)


def write_reference_sheet(
    workbook: xlsxwriter.Workbook, reference_rows: list[list[object]]
) -> None:
    worksheet = workbook.add_worksheet("reference_discharge_capacity")
    worksheet.freeze_panes(1, 0)
    worksheet.write_row(
        0,
        0,
        [
            "mat_step_index",
            "date",
            "charge_cycle_count_before_step",
            "capacity_ah",
            "sample_count",
        ],
    )
    for row_index, row in enumerate(reference_rows, start=1):
        worksheet.write_number(row_index, 0, int(row[0]))
        worksheet.write_string(row_index, 1, str(row[1]))
        worksheet.write_number(row_index, 2, int(row[2]))
        worksheet.write_number(row_index, 3, float(row[3]))
        worksheet.write_number(row_index, 4, int(row[4]))


def export_workbook(
    mat_path: Path,
    output_path: Path,
    capacity_by_cycle: list[float] | None,
) -> tuple[int, int]:
    mat_steps = load_mat_steps(mat_path)
    total_rows = sum(step.sample_count for step in mat_steps)
    total_sheets = max(1, math.ceil(total_rows / ROWS_PER_SHEET))
    reference_rows = compute_reference_discharge_rows(mat_path, mat_steps)

    workbook = xlsxwriter.Workbook(
        str(output_path),
        {"constant_memory": True, "strings_to_numbers": False},
    )
    workbook.use_zip64()

    readme = workbook.add_worksheet("README")
    write_readme(readme, mat_path.name, len(mat_steps), total_rows, capacity_by_cycle is not None)
    write_cycle_map_sheet(workbook, mat_steps)
    write_reference_sheet(workbook, reference_rows)

    row_iter = iter_data_rows(mat_steps, capacity_by_cycle)
    for sheet_number in range(1, total_sheets + 1):
        worksheet = workbook.add_worksheet(f"data_{sheet_number:02d}")
        worksheet.freeze_panes(1, 0)
        worksheet.write_row(0, 0, DATA_COLUMNS)
        rows_written = 0
        while rows_written < ROWS_PER_SHEET:
            try:
                row = next(row_iter)
            except StopIteration:
                break
            excel_row = rows_written + 1
            worksheet.write_string(excel_row, 0, row[0])
            worksheet.write_string(excel_row, 1, row[1])
            worksheet.write_string(excel_row, 2, row[2])
            worksheet.write_string(excel_row, 3, row[3])
            worksheet.write_number(excel_row, 4, row[4])
            if row[5] is not None:
                worksheet.write_string(excel_row, 5, row[5])
            rows_written += 1

    workbook.close()
    return len(mat_steps), total_rows


def get_sheet_row_counts(workbook_path: Path) -> tuple[list[str], int]:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        data_sheets = [name for name in workbook.sheetnames if name.startswith("data_")]
        total_rows = 0
        for sheet_name in data_sheets:
            total_rows += max(0, workbook[sheet_name].max_row - 1)
        return data_sheets, total_rows
    finally:
        workbook.close()


def main() -> None:
    root = Path.cwd()
    output_dir = root / OUTPUT_DIR_NAME
    output_dir.mkdir(exist_ok=True)

    mat_files = sorted(root.glob("RW*.mat"))
    if not mat_files:
        raise FileNotFoundError("No RW MAT files were found in the working directory.")

    rw9_pickle = root / RW9_PICKLE_NAME
    rw9_capacity_map: list[float] | None = None
    if rw9_pickle.exists():
        rw9_capacity_map = build_rw9_capacity_map(load_mat_steps(root / "RW9.mat"), rw9_pickle)

    for mat_path in mat_files:
        print(f"Exporting {mat_path.name}...")
        capacity_map = rw9_capacity_map if mat_path.stem.upper() == "RW9" and rw9_capacity_map else None
        output_path = output_dir / f"{mat_path.stem}_pkl_layout.xlsx"
        cycle_count, expected_rows = export_workbook(mat_path, output_path, capacity_map)
        data_sheets, actual_rows = get_sheet_row_counts(output_path)
        if actual_rows != expected_rows:
            raise ValueError(
                f"Workbook row-count check failed for {output_path.name}: {actual_rows} != {expected_rows}"
            )
        print(
            f"Verified {output_path.name}: cycles={cycle_count}, rows={actual_rows}, data_sheets={len(data_sheets)}"
        )

    print(f"Created pkl-layout Excel exports in: {output_dir}")


if __name__ == "__main__":
    main()
