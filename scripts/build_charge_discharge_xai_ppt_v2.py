from __future__ import annotations

import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(r"c:\Users\nimri\Downloads\BatteryData")
SUITE_DIR = ROOT / "charge_discharge_full_cycle_suite"
XAI_DIR = SUITE_DIR / "interpretability_shap_lime"
OUTDIR = ROOT / "ppt_charge_discharge_xai_summary_v2"
ASSET_DIR = OUTDIR / "assets"
PPT_PATH = OUTDIR / "charge_discharge_feature_token_transformer_xai_summary_v2.pptx"

BATTERY_COLORS = {"RW9": "#1f77b4", "RW10": "#ff7f0e", "RW11": "#2ca02c", "RW12": "#d62728"}
COLORS = {
    "navy": RGBColor(23, 43, 77),
    "muted": RGBColor(95, 102, 116),
    "blue": RGBColor(42, 108, 176),
    "orange": RGBColor(226, 116, 49),
    "green": RGBColor(47, 133, 90),
    "purple": RGBColor(107, 78, 161),
    "red": RGBColor(190, 75, 73),
    "light": RGBColor(245, 247, 250),
    "line": RGBColor(210, 216, 225),
    "white": RGBColor(255, 255, 255),
}


def ensure_dirs() -> None:
    OUTDIR.mkdir(exist_ok=True)
    ASSET_DIR.mkdir(exist_ok=True)


def parse_date(series: pd.Series) -> pd.Series:
    return pd.to_datetime(series, format="%d-%b-%Y %H:%M:%S", errors="coerce")


def load_reference_capacity(battery: str) -> pd.DataFrame:
    df = pd.read_csv(ROOT / "Analysis" / battery / "reference_capacity_summary.csv")
    df["date"] = parse_date(df["date"])
    out = (
        df.groupby("charge_cycle_count_before_reference", as_index=False)
        .agg(date=("date", "min"), capacity_ah=("capacity_ah", "mean"))
        .sort_values("date")
        .reset_index(drop=True)
    )
    initial = out.loc[out["charge_cycle_count_before_reference"].eq(0), "capacity_ah"].iloc[0]
    out["soh_percent"] = out["capacity_ah"] / initial * 100.0
    out["battery"] = battery
    return out


def save_capacity_soh_plot() -> Path:
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.0))
    for battery in ("RW9", "RW10", "RW11", "RW12"):
        df = load_reference_capacity(battery)
        axes[0].plot(df["date"], df["capacity_ah"], marker="o", ms=3.0, lw=1.6, label=battery, color=BATTERY_COLORS[battery])
        axes[1].plot(df["date"], df["soh_percent"], marker="o", ms=3.0, lw=1.6, label=battery, color=BATTERY_COLORS[battery])
    axes[0].set_title("Reference-discharge capacity")
    axes[0].set_ylabel("Capacity (Ah)")
    axes[1].set_title("SOH target")
    axes[1].set_ylabel("SOH (%)")
    for ax in axes:
        ax.grid(True, alpha=0.25)
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%d-%b"))
        for label in ax.get_xticklabels():
            label.set_rotation(30)
            label.set_ha("right")
    axes[0].legend(ncol=4, frameon=False, loc="upper right")
    fig.suptitle("Benchmark labels come from sparse reference-discharge checkpoints", fontsize=15, weight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.93))
    out = ASSET_DIR / "capacity_soh_reference_labels.png"
    fig.savefig(out, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return out


def save_event_coverage_plot() -> Path:
    event_summary = pd.read_csv(SUITE_DIR / "event_summary.csv")
    pivot = event_summary.pivot(index="battery", columns="event_type", values="events").loc[["RW9", "RW10", "RW11"]]
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 4.5))
    x = np.arange(len(pivot))
    width = 0.36
    axes[0].bar(x - width / 2, pivot["charge"], width, label="charge", color="#2a9d8f")
    axes[0].bar(x + width / 2, pivot["discharge"], width, label="discharge", color="#e76f51")
    axes[0].set_xticks(x, pivot.index)
    axes[0].set_ylabel("Events")
    axes[0].set_title("Charge/discharge event count")
    axes[0].legend(frameon=False)
    axes[0].grid(True, axis="y", alpha=0.25)

    rows = []
    for battery in ("RW9", "RW10", "RW11"):
        ref = load_reference_capacity(battery)
        rows.append({"battery": battery, "checkpoints": int((ref["charge_cycle_count_before_reference"] > 0).sum())})
    chk = pd.DataFrame(rows)
    axes[1].bar(chk["battery"], chk["checkpoints"], color="#457b9d")
    axes[1].set_ylabel("Reference SOH checkpoints")
    axes[1].set_title("Sparse independent SOH labels")
    axes[1].grid(True, axis="y", alpha=0.25)
    fig.suptitle("Discharge doubles events, but does not double independent SOH labels", fontsize=15, weight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.92))
    out = ASSET_DIR / "charge_discharge_event_coverage.png"
    fig.savefig(out, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return out


def save_top_shap_plot() -> Path:
    top = pd.read_csv(XAI_DIR / "feature_token_transformer" / "global_shap_feature_importance.csv").head(12).iloc[::-1]
    fig, ax = plt.subplots(figsize=(8.4, 5.0))
    ax.barh(top["feature"], top["mean_abs_shap_combined"], color="#6b4ea1")
    ax.set_xlabel("Mean |SHAP| in SOH percentage points")
    ax.set_title("Top global SHAP drivers")
    ax.grid(True, axis="x", alpha=0.25)
    fig.tight_layout()
    out = ASSET_DIR / "top_transformer_shap_features.png"
    fig.savefig(out, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return out


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.25), Inches(12.2), Inches(0.50))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = COLORS["navy"]
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.58), Inches(0.78), Inches(12.0), Inches(0.35))
        sp = sub.text_frame.paragraphs[0]
        sr = sp.add_run()
        sr.text = subtitle
        sr.font.size = Pt(11)
        sr.font.color.rgb = COLORS["muted"]


def add_footer(slide) -> None:
    box = slide.shapes.add_textbox(Inches(0.55), Inches(7.12), Inches(12.2), Inches(0.20))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = "NASA RW battery SOH modeling | charge + discharge | feature-token transformer"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(135, 142, 154)


def fit_image(slide, path: Path, x: float, y: float, w: float, h: float) -> None:
    with Image.open(path) as image:
        img_w, img_h = image.size
    img_ratio = img_w / img_h
    box_ratio = w / h
    if img_ratio >= box_ratio:
        final_w = w
        final_h = w / img_ratio
    else:
        final_h = h
        final_w = h * img_ratio
    final_x = x + (w - final_w) / 2.0
    final_y = y + (h - final_h) / 2.0
    slide.shapes.add_picture(str(path), Inches(final_x), Inches(final_y), width=Inches(final_w), height=Inches(final_h))


def add_textbox(slide, x: float, y: float, w: float, h: float, text: str, size: int = 12, bold: bool = False, color=COLORS["navy"]) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_bullets(slide, x: float, y: float, w: float, h: float, bullets: list[str], size: int = 12) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        p.font.size = Pt(size)
        p.font.color.rgb = COLORS["navy"]
        p.space_after = Pt(6)


def add_table(slide, df: pd.DataFrame, x: float, y: float, w: float, h: float, font_size: int = 8) -> None:
    table = slide.shapes.add_table(df.shape[0] + 1, df.shape[1], Inches(x), Inches(y), Inches(w), Inches(h)).table
    for col_idx, col in enumerate(df.columns):
        cell = table.cell(0, col_idx)
        cell.text = str(col)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLORS["navy"]
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(font_size)
            p.font.bold = True
            p.font.color.rgb = COLORS["white"]
    for row_pos, (_, row) in enumerate(df.iterrows()):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_pos + 1, col_idx)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(250, 251, 253) if row_pos % 2 == 0 else RGBColor(238, 242, 247)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(font_size)
                p.font.color.rgb = COLORS["navy"]


def add_block(slide, x: float, y: float, w: float, h: float, text: str, fill: RGBColor) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = COLORS["line"]
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = COLORS["white"]


def add_arrow(slide, x1: float, y1: float, x2: float, y2: float) -> None:
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = RGBColor(120, 128, 140)
    line.line.width = Pt(1.4)
    line.line.end_arrowhead = True


def build_deck() -> None:
    ensure_dirs()
    capacity_plot = save_capacity_soh_plot()
    coverage_plot = save_event_coverage_plot()
    top_shap_plot = save_top_shap_plot()

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    best_metrics = json.loads((SUITE_DIR / "feature_tokens_h4_d64_l2" / "metrics.json").read_text())
    local_instances = pd.read_csv(XAI_DIR / "local_instances_used.csv")
    top_shap = pd.read_csv(XAI_DIR / "feature_token_transformer" / "global_shap_feature_importance.csv").head(8)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Charge + Discharge SOH Prediction With XAI", "Revised deck: figures preserve aspect ratio and slides are separated to avoid overlap.")
    add_textbox(slide, 0.9, 1.45, 5.6, 0.35, "Best charge+discharge model", 18, True, COLORS["purple"])
    add_bullets(slide, 0.9, 2.0, 5.8, 2.4, [
        "Feature-token pure transformer.",
        "Train: RW9 + RW10 charge/discharge events.",
        "Test: RW11 charge/discharge events.",
        "Target: next reference-discharge SOH checkpoint.",
        "RW11 R² = 0.9331, MAE = 3.20, RMSE = 3.71.",
    ], 15)
    add_textbox(slide, 7.3, 1.45, 4.9, 0.35, "XAI included", 18, True, COLORS["green"])
    add_bullets(slide, 7.3, 2.0, 5.0, 2.4, [
        "Global SHAP for combined, charge, and discharge events.",
        "Local LIME for one representative charge event.",
        "Local LIME for one representative discharge event.",
        "Interpretation slide explaining time/cycle proxy features.",
    ], 15)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Why Time and Count Features Affect Prediction", "These are not the target label, but they are strong proxies for battery aging.")
    add_bullets(slide, 0.75, 1.25, 5.8, 4.9, [
        "time_start and time_end are absolute experiment time, not relTime.",
        "label_charge_count is the number of charge random-walk events completed before the current event.",
        "As aging progresses, capacity/SOH generally decreases, so chronological features are highly correlated with the target.",
        "sample_count and duration can also become aging proxies because event length changes as the cell behavior changes under the randomized protocol.",
        "If the goal is waveform-only SOH estimation, remove time_start, time_end, and label_charge_count and report that result separately.",
    ], 15)
    proxy_df = pd.DataFrame([
        ["relTime", "Within-event time axis", "Shape/duration of one event"],
        ["time_start/end", "Absolute experiment clock", "Battery age / calendar position"],
        ["label_charge_count", "Cumulative charge-event count", "Aging progression proxy"],
        ["sample_count", "Rows in event", "Protocol response / duration proxy"],
    ], columns=["Feature", "Meaning", "Why it matters"])
    add_table(slide, proxy_df, 6.85, 1.45, 5.7, 2.6, 8)
    add_textbox(slide, 6.9, 4.55, 5.6, 1.0, "Important caveat: SHAP ranking these features highly means the model is using aging chronology strongly. This can be valid only if time/history is available at inference.", 13, True, COLORS["red"])
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Reference Capacity and SOH Labels", "Capacity labels are measured at sparse reference-discharge checkpoints.")
    fit_image(slide, capacity_plot, 0.65, 1.15, 12.0, 5.55)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Charge + Discharge Event Coverage", "Adding discharge increases event count, but labels remain checkpoint-level.")
    fit_image(slide, coverage_plot, 0.85, 1.2, 11.65, 4.9)
    coverage_df = pd.DataFrame([
        ["Train", "RW9 + RW10", "94,965 events"],
        ["Test", "RW11", "46,490 events"],
        ["RW11 labels", "Reference checkpoints", "38 benchmark SOH values"],
    ], columns=["Split", "Source", "Count"])
    add_table(slide, coverage_df, 2.1, 6.12, 9.1, 0.8, 9)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Feature-Token Transformer Workflow", "One event becomes 63 engineered features, grouped into 8 transformer tokens.")
    y = 1.55
    add_block(slide, 0.55, y, 1.55, 0.72, "Raw MAT\nsteps", COLORS["muted"])
    add_arrow(slide, 2.1, y + 0.36, 2.55, y + 0.36)
    add_block(slide, 2.55, y, 1.75, 0.72, "Charge +\ndischarge events", COLORS["blue"])
    add_arrow(slide, 4.3, y + 0.36, 4.75, y + 0.36)
    add_block(slide, 4.75, y, 1.75, 0.72, "Next reference\nSOH label", COLORS["green"])
    add_arrow(slide, 6.5, y + 0.36, 6.95, y + 0.36)
    add_block(slide, 6.95, y, 1.75, 0.72, "63 engineered\nfeatures", COLORS["orange"])
    add_arrow(slide, 8.7, y + 0.36, 9.15, y + 0.36)
    add_block(slide, 9.15, y, 1.55, 0.72, "8 feature\ntokens", COLORS["purple"])
    add_arrow(slide, 10.7, y + 0.36, 11.15, y + 0.36)
    add_block(slide, 11.15, y, 1.45, 0.72, "SOH\nprediction", COLORS["red"])
    add_bullets(slide, 0.8, 3.0, 5.8, 2.4, [
        "No separate MLP feature branch.",
        "The engineered feature blocks are the transformer tokens.",
        "CLS output is passed to a regression head.",
        "Loss uses inverse benchmark-block weighted MSE.",
    ], 13)
    params = pd.DataFrame([
        ["Features", "63"],
        ["Feature tokens", "8"],
        ["d_model", best_metrics["spec"]["d_model"]],
        ["Heads", best_metrics["spec"]["nhead"]],
        ["Layers", best_metrics["spec"]["num_layers"]],
        ["Dropout", best_metrics["spec"]["dropout"]],
        ["Learning rate", best_metrics["spec"]["learning_rate"]],
    ], columns=["Parameter", "Value"])
    add_table(slide, params, 7.2, 3.0, 4.8, 2.55, 9)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Model Results", "Charge-only and charge+discharge experiments side by side.")
    fit_image(slide, SUITE_DIR / "combined_charge_vs_charge_discharge_summary_with_table.png", 0.45, 1.0, 12.45, 6.0)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Best Transformer Training Curves", "Feature-token pure transformer, charge + discharge.")
    fit_image(slide, SUITE_DIR / "feature_tokens_h4_d64_l2" / "training_dashboard.png", 0.9, 1.18, 11.45, 5.55)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Best Transformer Prediction Report", "RW11 actual vs predicted SOH and checkpoint aggregation.")
    fit_image(slide, SUITE_DIR / "feature_tokens_h4_d64_l2" / "prediction_report.png", 0.75, 1.05, 11.85, 5.9)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Global SHAP: Charge vs Discharge Drivers", "Mean absolute SHAP values by event type.")
    fit_image(slide, XAI_DIR / "feature_token_transformer" / "global_shap_bar_by_event_type.png", 0.9, 1.1, 11.3, 5.65)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Top Global SHAP Features", "The transformer relies strongly on aging-chronology proxies.")
    fit_image(slide, top_shap_plot, 1.0, 1.15, 6.4, 5.2)
    shap_df = top_shap[["feature", "mean_abs_shap_combined"]].copy()
    shap_df["mean_abs_shap_combined"] = shap_df["mean_abs_shap_combined"].map(lambda x: f"{x:.3f}")
    shap_df.columns = ["Feature", "Mean |SHAP|"]
    add_table(slide, shap_df, 7.7, 1.28, 4.5, 3.45, 8)
    add_textbox(slide, 7.8, 5.05, 4.4, 0.95, "Interpretation: time_start, time_end, and label_charge_count encode where the event sits in the aging timeline.", 12, True, COLORS["red"])
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "SHAP Beeswarm: Charge Events", "Feature effects for sampled RW11 charge events.")
    fit_image(slide, XAI_DIR / "feature_token_transformer" / "global_shap_beeswarm_charge.png", 0.95, 1.0, 11.3, 6.0)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "SHAP Beeswarm: Discharge Events", "Feature effects for sampled RW11 discharge events.")
    fit_image(slide, XAI_DIR / "feature_token_transformer" / "global_shap_beeswarm_discharge.png", 0.95, 1.0, 11.3, 6.0)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Local LIME: Representative Charge Event", "Local explanation for one RW11 charge random-walk event.")
    fit_image(slide, XAI_DIR / "feature_token_transformer" / "local_lime_charge.png", 0.95, 1.05, 11.2, 5.35)
    charge = local_instances[local_instances["event_type"].eq("charge")].copy()
    charge["actual_soh_percent"] = charge["actual_soh_percent"].map(lambda x: f"{x:.2f}")
    charge["transformer_prediction"] = charge["transformer_prediction"].map(lambda x: f"{x:.2f}")
    add_table(slide, charge[["event_type", "event_order", "assigned_checkpoint_cycle", "actual_soh_percent", "transformer_prediction"]], 1.8, 6.45, 9.7, 0.45, 8)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Local LIME: Representative Discharge Event", "Local explanation for one RW11 discharge random-walk event.")
    fit_image(slide, XAI_DIR / "feature_token_transformer" / "local_lime_discharge.png", 0.95, 1.05, 11.2, 5.35)
    discharge = local_instances[local_instances["event_type"].eq("discharge")].copy()
    discharge["actual_soh_percent"] = discharge["actual_soh_percent"].map(lambda x: f"{x:.2f}")
    discharge["transformer_prediction"] = discharge["transformer_prediction"].map(lambda x: f"{x:.2f}")
    add_table(slide, discharge[["event_type", "event_order", "assigned_checkpoint_cycle", "actual_soh_percent", "transformer_prediction"]], 1.8, 6.45, 9.7, 0.45, 8)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_title(slide, "Takeaways and Scientific Caveat", "What the current result means.")
    add_bullets(slide, 0.85, 1.25, 11.8, 4.8, [
        "Best charge+discharge transformer: feature-token pure transformer, R² = 0.9331 on RW11.",
        "The strongest XAI drivers are chronology proxies: time_start, time_end, and label_charge_count.",
        "sample_count and duration also matter because event length changes with operating response and aging.",
        "This is valid if inference has experiment history or cycle/time information.",
        "For a waveform-only claim, rerun after removing absolute time and label_charge_count.",
    ], 16)
    add_footer(slide)

    prs.save(PPT_PATH)
    print(PPT_PATH)


if __name__ == "__main__":
    build_deck()
