from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUTDIR = ROOT / "ppt_battery_transformer_summary"
ASSET_DIR = OUTDIR / "assets"
PPT_PATH = OUTDIR / "battery_capacity_transformer_summary.pptx"

BATTERIES = ("RW9", "RW10", "RW11", "RW12")
TRAIN_BATTERIES = ("RW9", "RW10")
TEST_BATTERY = "RW11"

COLORS = {
    "RW9": "#1f77b4",
    "RW10": "#ff7f0e",
    "RW11": "#2ca02c",
    "RW12": "#d62728",
}


@dataclass(frozen=True)
class EngineeredFeatureRow:
    family: str
    derived_from: str
    example_columns: str
    interpretation: str


ENGINEERED_FEATURE_ROWS = [
    EngineeredFeatureRow(
        "Charge throughput",
        "current + time",
        "charge_throughput_ah",
        "Charge moved in the segment",
    ),
    EngineeredFeatureRow(
        "Energy throughput",
        "voltage + current + time",
        "energy_throughput_wh, power_mean_w",
        "Electrical work / usage severity",
    ),
    EngineeredFeatureRow(
        "Voltage shape",
        "voltage + time",
        "voltage_start, voltage_end, voltage_delta, dv_dt_*",
        "Charge progress and electrochemical response",
    ),
    EngineeredFeatureRow(
        "Current statistics",
        "current",
        "current_mean_a, current_abs_mean_a, current_std_a",
        "Loading severity under randomized charging",
    ),
    EngineeredFeatureRow(
        "Thermal response",
        "temperature + time",
        "temperature_mean_c, temperature_delta_c, dtemp_dt_*",
        "Thermal stress and heating dynamics",
    ),
    EngineeredFeatureRow(
        "Voltage-window timing",
        "voltage + time",
        "time_in_vwin_*, time_cross_3p9_to_4p1_s",
        "How quickly comparable voltage regions are traversed",
    ),
    EngineeredFeatureRow(
        "IC / DV features",
        "voltage + charge",
        "dq_dv_*, dv_dq_*",
        "Aging-sensitive shape changes",
    ),
    EngineeredFeatureRow(
        "Distributional statistics",
        "voltage / temperature sequences",
        "voltage_pdf_*, temperature_pdf_*",
        "Sequence-level distribution behavior",
    ),
]


def parse_datetime(series: pd.Series) -> pd.Series:
    return pd.to_datetime(series, format="%d-%b-%Y %H:%M:%S", errors="coerce")


def ensure_dirs() -> None:
    OUTDIR.mkdir(exist_ok=True)
    ASSET_DIR.mkdir(exist_ok=True)


def load_reference_capacity_summary(battery: str) -> pd.DataFrame:
    path = ROOT / "Analysis" / battery / "reference_capacity_summary.csv"
    df = pd.read_csv(path)
    df["date"] = parse_datetime(df["date"])
    grouped = (
        df.groupby("charge_cycle_count_before_reference", as_index=False)
        .agg(
            date=("date", "min"),
            capacity_ah=("capacity_ah", "mean"),
            sample_count=("sample_count", "sum"),
            n_raw_reference_steps=("mat_step_index", "count"),
        )
        .sort_values("date")
        .reset_index(drop=True)
    )
    init_capacity = grouped.loc[
        grouped["charge_cycle_count_before_reference"].eq(0), "capacity_ah"
    ].iloc[0]
    grouped["soh_percent"] = 100.0 * grouped["capacity_ah"] / init_capacity
    grouped["battery"] = battery
    return grouped


def load_charge_features(battery: str, columns: Iterable[str]) -> pd.DataFrame:
    path = ROOT / "charge engineered features" / f"{battery}_charge_engineered_features.csv"
    df = pd.read_csv(path, usecols=list(columns))
    if "date" in df.columns:
        df["date"] = parse_datetime(df["date"])
    return df


def save_capacity_over_time_plot() -> Path:
    fig, axes = plt.subplots(2, 2, figsize=(14, 8), sharex=False, sharey=False)
    for ax, battery in zip(axes.flat, BATTERIES):
        df = load_reference_capacity_summary(battery)
        ax.plot(
            df["date"],
            df["capacity_ah"],
            marker="o",
            ms=3.5,
            lw=1.8,
            color=COLORS[battery],
        )
        ax.set_title(f"{battery}: unique reference checkpoints", fontsize=11, weight="bold")
        ax.set_ylabel("Capacity (Ah)")
        ax.grid(True, alpha=0.25)
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%d-%b"))
        for label in ax.get_xticklabels():
            label.set_rotation(30)
            label.set_ha("right")
    fig.suptitle("Reference-discharge capacity over calendar time", fontsize=16, weight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    out = ASSET_DIR / "capacity_over_time_by_battery.png"
    fig.savefig(out, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return out


def _daily_feature_summary(feature_names: list[str]) -> pd.DataFrame:
    frames: list[pd.DataFrame] = []
    usecols = ["battery", "date"] + feature_names
    for battery in BATTERIES:
        df = load_charge_features(battery, usecols)
        df = df.dropna(subset=["date"]).sort_values("date")
        daily = (
            df.set_index("date")[feature_names]
            .resample("3D")
            .median()
            .interpolate(limit_direction="both")
        )
        daily["battery"] = battery
        daily = daily.reset_index()
        frames.append(daily)
    return pd.concat(frames, ignore_index=True)


def save_raw_input_over_time_plot() -> Path:
    features = [
        ("voltage_start", "Voltage start (V)"),
        ("current_mean_a", "Mean current (A)"),
        ("temperature_mean_c", "Mean temperature (C)"),
        ("duration_s", "Cycle duration (s)"),
    ]
    df = _daily_feature_summary([name for name, _ in features])
    fig, axes = plt.subplots(2, 2, figsize=(14, 8), sharex=True)
    for ax, (feature, label) in zip(axes.flat, features):
        for battery in BATTERIES:
            sub = df[df["battery"] == battery]
            ax.plot(sub["date"], sub[feature], lw=1.7, label=battery, color=COLORS[battery])
        ax.set_title(label, fontsize=11, weight="bold")
        ax.grid(True, alpha=0.25)
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%d-%b"))
    axes[0, 0].legend(ncol=4, loc="upper center", bbox_to_anchor=(1.05, 1.35), frameon=False)
    fig.suptitle(
        "Cycle-level summaries of raw charge inputs over calendar time",
        fontsize=16,
        weight="bold",
    )
    fig.tight_layout(rect=(0, 0, 1, 0.95))
    out = ASSET_DIR / "raw_inputs_over_time.png"
    fig.savefig(out, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return out


def save_engineered_feature_over_time_plot() -> Path:
    features = [
        ("charge_throughput_ah", "Charge throughput (Ah)"),
        ("energy_throughput_wh", "Energy throughput (Wh)"),
        ("voltage_delta", "Voltage delta (V)"),
        ("dq_dv_median_ahpv", "Median dQ/dV"),
    ]
    df = _daily_feature_summary([name for name, _ in features])
    fig, axes = plt.subplots(2, 2, figsize=(14, 8), sharex=True)
    for ax, (feature, label) in zip(axes.flat, features):
        for battery in BATTERIES:
            sub = df[df["battery"] == battery]
            ax.plot(sub["date"], sub[feature], lw=1.7, label=battery, color=COLORS[battery])
        ax.set_title(label, fontsize=11, weight="bold")
        ax.grid(True, alpha=0.25)
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%d-%b"))
    axes[0, 0].legend(ncol=4, loc="upper center", bbox_to_anchor=(1.05, 1.35), frameon=False)
    fig.suptitle(
        "Engineered charge features over calendar time (3-day median)",
        fontsize=16,
        weight="bold",
    )
    fig.tight_layout(rect=(0, 0, 1, 0.95))
    out = ASSET_DIR / "engineered_features_over_time.png"
    fig.savefig(out, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return out


def add_title(slide, text: str, subtitle: str | None = None) -> None:
    title_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.28), Inches(12.2), Inches(0.6))
    p = title_box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(23, 43, 77)
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.83), Inches(12.0), Inches(0.35))
        sp = sub_box.text_frame.paragraphs[0]
        srun = sp.add_run()
        srun.text = subtitle
        srun.font.size = Pt(11)
        srun.font.color.rgb = RGBColor(90, 96, 110)


def add_footer(slide, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.6), Inches(6.95), Inches(12.0), Inches(0.25))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = text
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(120, 120, 120)


def add_bullets(slide, items: list[str], left: float, top: float, width: float, height: float) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(15)
        p.font.color.rgb = RGBColor(35, 35, 35)
        p.space_after = Pt(6)


def add_table(
    slide,
    data: list[list[str]],
    left: float,
    top: float,
    width: float,
    height: float,
    col_widths: list[float] | None = None,
    font_size: int = 11,
) -> None:
    rows = len(data)
    cols = len(data[0])
    table = slide.shapes.add_table(rows, cols, Inches(left), Inches(top), Inches(width), Inches(height)).table
    if col_widths:
        for i, col_w in enumerate(col_widths):
            table.columns[i].width = Inches(col_w)
    for r, row in enumerate(data):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.margin_left = Inches(0.04)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(font_size)
            p.font.color.rgb = RGBColor(25, 25, 25)
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(225, 232, 245)
                p.font.bold = True
            else:
                if r % 2 == 1:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(247, 249, 252)


def build_presentation(
    capacity_plot: Path,
    raw_inputs_plot: Path,
    engineered_plot: Path,
) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]

    metrics = json.loads(
        (ROOT / "cycle_level_weighted_transformer_rw9_rw10_to_rw11" / "metrics.json").read_text()
    )
    dataset_summary = pd.read_csv(
        ROOT / "cycle_level_weighted_transformer_rw9_rw10_to_rw11" / "dataset_summary.csv"
    )
    battery_summary = pd.read_csv(ROOT / "Analysis" / "aggregate" / "battery_comparison_summary.csv")
    driver_scores = pd.read_csv(ROOT / "rw12_driver_ranges" / "rw11_rw12_matched_soh_driver_scores.csv").head(5)

    # Slide 1
    slide = prs.slides.add_slide(blank)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(22, 54, 92)
    shape.line.color.rgb = RGBColor(22, 54, 92)
    title = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(9.8), Inches(0.75))
    p = title.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "Battery Capacity, Feature Analysis, and Transformer Results"
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
    subtitle = slide.shapes.add_textbox(Inches(0.75), Inches(1.75), Inches(11.0), Inches(1.5))
    s = subtitle.text_frame
    p1 = s.paragraphs[0]
    p1.text = "Dataset focus: NASA Randomized Battery Usage (RW9, RW10, RW11, RW12)"
    p1.font.size = Pt(19)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(23, 43, 77)
    p2 = s.add_paragraph()
    p2.text = "Model focus: cycle-level weighted transformer trained on RW9+RW10 and tested on RW11; RW12 analyzed separately for shift behavior."
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(70, 70, 70)
    add_footer(slide, "Generated from local BatteryData analyses")

    # Slide 2
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Capacity Over Time for Each Cell")
    slide.shapes.add_picture(str(capacity_plot), Inches(0.55), Inches(1.15), width=Inches(8.5))
    table_data = [["Battery", "Charge cycles", "Reference steps", "Initial Ah", "Final Ah"]]
    for _, row in battery_summary.iterrows():
        table_data.append(
            [
                str(row["battery"]),
                f"{int(row['charge_segments']):,}",
                f"{int(row['reference_count'])}",
                f"{row['initial_capacity_ah']:.3f}",
                f"{row['final_capacity_ah']:.3f}",
            ]
        )
    add_table(slide, table_data, 9.2, 1.55, 3.45, 2.65, [0.75, 1.0, 0.8, 0.45, 0.45], 10)
    add_bullets(
        slide,
        [
            "Each point is a unique benchmark reference-discharge checkpoint after grouping duplicate raw discharge steps by the same checkpoint index.",
            "RW12 degrades more slowly and stays at higher capacity later in calendar time than RW9-RW11.",
        ],
        9.2,
        4.45,
        3.5,
        1.75,
    )
    add_footer(slide, "Capacity labels used later for SOH come from these reference-discharge checkpoints")

    # Slide 3
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Raw Charge Inputs Over Time")
    slide.shapes.add_picture(str(raw_inputs_plot), Inches(0.55), Inches(1.15), width=Inches(12.15))
    add_footer(slide, "Cycle-level medians aggregated every 3 days; Current < 0 indicates charging in the charge-only files")

    # Slide 4
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Engineered Feature Inventory")
    inventory_data = [
        ["Family", "Derived from", "Example columns", "What it measures"],
    ]
    for row in ENGINEERED_FEATURE_ROWS:
        inventory_data.append([row.family, row.derived_from, row.example_columns, row.interpretation])
    add_table(
        slide,
        inventory_data,
        0.45,
        1.2,
        12.35,
        5.6,
        [1.65, 1.65, 3.7, 5.2],
        10,
    )
    add_footer(slide, "Feature families extracted from the charge segments after temperature alignment from the raw MATLAB files")

    # Slide 5
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Engineered Feature Trends Over Time")
    slide.shapes.add_picture(str(engineered_plot), Inches(0.55), Inches(1.15), width=Inches(12.15))
    add_footer(slide, "These features were extracted per charge cycle and later reused for modeling and shift analysis")

    # Slide 6
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Driver Analysis: Why RW12 Behaves Differently")
    slide.shapes.add_picture(
        str(ROOT / "rw12_driver_ranges" / "02_driver_strength_bar.png"),
        Inches(0.55),
        Inches(1.2),
        width=Inches(5.2),
    )
    slide.shapes.add_picture(
        str(ROOT / "rw12_driver_ranges" / "03_matched_soh_feature_profiles.png"),
        Inches(5.95),
        Inches(1.2),
        width=Inches(6.8),
    )
    bullets = [
        "RW12 mostly stays inside the observed feature ranges, so failure is not simple out-of-range extrapolation.",
        "The stronger issue is conditional mismatch: at the same SOH, RW12 tends to start charging from higher voltage, move less charge/energy, and traverse a smaller voltage window.",
    ]
    add_bullets(slide, bullets, 0.7, 5.95, 12.0, 0.75)
    top_driver_text = "Top matched-SOH drivers: " + ", ".join(
        f"{row.feature} ({row.avg_abs_effect_size:.3f})" for row in driver_scores.itertuples(index=False)
    )
    extra = slide.shapes.add_textbox(Inches(0.7), Inches(6.55), Inches(12.0), Inches(0.35))
    p = extra.text_frame.paragraphs[0]
    p.text = top_driver_text
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(60, 60, 60)
    add_footer(slide, "Driver scores compare RW11 and RW12 after matching by SOH bins")

    # Slide 7
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Cycle-Level Weighted Transformer")
    slide.shapes.add_picture(
        str(ROOT / "cycle_level_weighted_transformer_rw9_rw10_to_rw11" / "architecture_diagram.png"),
        Inches(0.55),
        Inches(1.2),
        width=Inches(8.0),
    )
    add_bullets(
        slide,
        [
            "One sample = one full charge (random walk) cycle.",
            "Each cycle is split into contiguous 30-row patches, giving up to 11 tokens per cycle.",
            "Each token contains 26 patch-level features built from voltage, current, temperature, and time.",
            "Target = SOH (%) of the next benchmark reference-discharge checkpoint.",
            "Training uses inverse checkpoint-block weighting so repeated labels do not dominate the loss.",
        ],
        8.75,
        1.45,
        3.75,
        4.8,
    )
    add_footer(slide, "This is the best cross-battery run completed so far: train RW9+RW10, test RW11")

    # Slide 8
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Model Settings and Dataset Split")
    spec = metrics["spec"]
    hyper_table = [
        ["Parameter", "Value"],
        ["Model", "Pure transformer encoder"],
        ["d_model", str(spec["d_model"])],
        ["Attention heads", str(spec["nhead"])],
        ["Encoder layers", str(spec["num_layers"])],
        ["Feedforward dim", str(spec["dim_feedforward"])],
        ["Dropout", f"{spec['dropout']:.2f}"],
        ["Learning rate", f"{spec['learning_rate']:.1e}"],
        ["Weight decay", f"{spec['weight_decay']:.1e}"],
        ["Batch size", str(spec["batch_size"])],
        ["Epochs", str(spec["epochs"])],
        ["Max tokens / cycle", str(spec["max_tokens"])],
        ["Input dim / token", str(metrics["input_dim"])],
        ["Trainable parameters", f"{metrics['parameter_count']:,}"],
    ]
    add_table(slide, hyper_table, 0.7, 1.35, 4.1, 5.5, [2.1, 1.8], 11)

    split_table = [["Split", "Battery set", "Cycles", "Unique checkpoints", "Median tokens"]]
    for _, row in dataset_summary.iterrows():
        battery_set = "RW9 + RW10" if row["split"] == "train" else "RW11"
        split_table.append(
            [
                str(row["split"]).title(),
                battery_set,
                f"{int(row['sample_count']):,}",
                f"{int(row['unique_checkpoint_labels'])}",
                f"{row['median_token_count']:.0f}",
            ]
        )
    add_table(slide, split_table, 5.25, 1.35, 4.8, 1.8, [1.0, 1.8, 1.2, 1.4, 1.1], 11)

    add_bullets(
        slide,
        [
            "Loss: weighted MSE with sample weight = 1 / benchmark-block size.",
            "Feature scaling: train-set standardization applied to token features and target SOH.",
            "Padding mask used for short cycles; CLS token used for regression output.",
        ],
        5.35,
        3.55,
        6.8,
        1.7,
    )
    add_footer(slide, "RW12 is completely excluded from this model run")

    # Slide 9
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Training Dynamics and Aggregate Performance")
    slide.shapes.add_picture(
        str(ROOT / "cycle_level_weighted_transformer_rw9_rw10_to_rw11" / "training_dashboard.png"),
        Inches(0.55),
        Inches(1.25),
        width=Inches(7.3),
    )
    result_table = [
        ["Evaluation", "R²", "MAE", "RMSE", "MAPE"],
        [
            "Train, cycle-level",
            f"{metrics['train_metrics_cycle_level']['r2']:.4f}",
            f"{metrics['train_metrics_cycle_level']['mae']:.3f}",
            f"{metrics['train_metrics_cycle_level']['rmse']:.3f}",
            f"{metrics['train_metrics_cycle_level']['mape_percent']:.2f}%",
        ],
        [
            "Test, cycle-level",
            f"{metrics['test_metrics_cycle_level']['r2']:.4f}",
            f"{metrics['test_metrics_cycle_level']['mae']:.3f}",
            f"{metrics['test_metrics_cycle_level']['rmse']:.3f}",
            f"{metrics['test_metrics_cycle_level']['mape_percent']:.2f}%",
        ],
        [
            "Test, checkpoint-avg",
            f"{metrics['test_metrics_checkpoint_aggregated']['r2']:.4f}",
            f"{metrics['test_metrics_checkpoint_aggregated']['mae']:.3f}",
            f"{metrics['test_metrics_checkpoint_aggregated']['rmse']:.3f}",
            f"{metrics['test_metrics_checkpoint_aggregated']['mape_percent']:.2f}%",
        ],
    ]
    add_table(slide, result_table, 8.15, 1.55, 4.55, 1.8, [1.8, 0.7, 0.7, 0.7, 0.65], 11)
    add_bullets(
        slide,
        [
            "Best observed RW11 test epoch during training: epoch 8 with R² = 0.9297 and MAE = 3.240.",
            "Cycle-level and checkpoint-aggregated metrics are both reported because many cycles share the same benchmark label.",
        ],
        8.15,
        3.75,
        4.4,
        2.2,
    )
    add_footer(slide, "Regression task: SOH (%) prediction, not classification")

    # Slide 10
    slide = prs.slides.add_slide(blank)
    add_title(slide, "RW11 Test Predictions")
    slide.shapes.add_picture(
        str(ROOT / "cycle_level_weighted_transformer_rw9_rw10_to_rw11" / "test_cycle_report.png"),
        Inches(0.45),
        Inches(1.15),
        width=Inches(6.25),
    )
    slide.shapes.add_picture(
        str(ROOT / "cycle_level_weighted_transformer_rw9_rw10_to_rw11" / "test_checkpoint_report.png"),
        Inches(6.95),
        Inches(1.15),
        width=Inches(5.95),
    )
    add_footer(slide, "Left: cycle-level predictions on all RW11 charge cycles. Right: predictions averaged back to benchmark checkpoints.")

    prs.save(PPT_PATH)


def main() -> None:
    ensure_dirs()
    capacity_plot = save_capacity_over_time_plot()
    raw_inputs_plot = save_raw_input_over_time_plot()
    engineered_plot = save_engineered_feature_over_time_plot()
    build_presentation(capacity_plot, raw_inputs_plot, engineered_plot)
    print(PPT_PATH)


if __name__ == "__main__":
    main()
