from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


SOURCE_DIR_NAME = "excel files version"
OUTPUT_DIR_NAME = "excel files version readable"
NUMERIC_FIELDS = ["time", "relativeTime", "voltage", "current", "temperature"]


def build_field_formula(field_name: str, end_column_letter: str) -> str:
    return (
        f'=LET(n,$B$5,'
        f'm,MATCH($B$1,{field_name}!$A:$A,0),'
        f'chunkCount,ROUNDUP(n/1000,0),'
        f'chunks,INDEX({field_name}!$C:${end_column_letter},m,SEQUENCE(1,chunkCount)),'
        f'vals,TOCOL(TEXTSPLIT(TEXTJOIN("|",TRUE,chunks),"|"),1),'
        f'TAKE(vals,n))'
    )


def add_step_view(workbook: openpyxl.Workbook) -> None:
    if "step_view" in workbook.sheetnames:
        del workbook["step_view"]

    validation_index = workbook.sheetnames.index("validation")
    worksheet = workbook.create_sheet("step_view", validation_index)
    worksheet.freeze_panes = "A8"
    worksheet.column_dimensions["A"].width = 18
    for column_letter in ["B", "C", "D", "E", "F", "G"]:
        worksheet.column_dimensions[column_letter].width = 22

    total_steps = workbook["step_index"].max_row - 1

    worksheet["A1"] = "selected_step_index"
    worksheet["B1"] = 1
    validation = DataValidation(type="whole", operator="between", formula1="1", formula2=str(total_steps))
    worksheet.add_data_validation(validation)
    validation.add(worksheet["B1"])

    worksheet["A2"] = "comment"
    worksheet["B2"] = '=XLOOKUP($B$1,step_index!$A:$A,step_index!$B:$B,"")'
    worksheet["A3"] = "type"
    worksheet["B3"] = '=XLOOKUP($B$1,step_index!$A:$A,step_index!$C:$C,"")'
    worksheet["A4"] = "date"
    worksheet["B4"] = '=XLOOKUP($B$1,step_index!$A:$A,step_index!$D:$D,"")'
    worksheet["A5"] = "sample_count"
    worksheet["B5"] = '=XLOOKUP($B$1,step_index!$A:$A,step_index!$E:$E,0)'
    worksheet["A6"] = "note"
    worksheet["B6"] = "Change B1 to any step index. The table below spills the exact text values for that MATLAB step."

    worksheet.append(["sample_index", "time", "relativeTime", "voltage", "current", "temperature"])
    worksheet["A8"] = "=SEQUENCE($B$5)"

    for column_index, field_name in enumerate(NUMERIC_FIELDS, start=2):
        source_sheet = workbook[field_name]
        end_column_letter = get_column_letter(source_sheet.max_column)
        worksheet.cell(row=8, column=column_index).value = build_field_formula(field_name, end_column_letter)


def maybe_update_readme(workbook: openpyxl.Workbook) -> None:
    if "README" not in workbook.sheetnames:
        return
    worksheet = workbook["README"]
    notes = [worksheet.cell(row=row_index, column=1).value for row_index in range(1, worksheet.max_row + 1)]
    extra_lines = [
        "Use the step_view sheet to reconstruct one MATLAB step as a normal table.",
        "In step_view!B1, enter the step_index you want to inspect.",
    ]
    next_row = worksheet.max_row + 1
    for line in extra_lines:
        if line not in notes:
            worksheet.cell(row=next_row, column=1).value = line
            next_row += 1


def main() -> None:
    root = Path.cwd()
    source_dir = root / SOURCE_DIR_NAME
    if not source_dir.exists():
        raise FileNotFoundError(f"Missing source folder: {source_dir}")

    output_dir = root / OUTPUT_DIR_NAME
    output_dir.mkdir(exist_ok=True)

    for source_path in sorted(source_dir.glob("*.xlsx")):
        print(f"Adding step_view to {source_path.name}...")
        workbook = openpyxl.load_workbook(source_path)
        add_step_view(workbook)
        maybe_update_readme(workbook)
        workbook.save(output_dir / source_path.name)
        workbook.close()
        print(f"Created {source_path.name}")

    print(f"Readable workbooks created in: {output_dir}")


if __name__ == "__main__":
    main()
